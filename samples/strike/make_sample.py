"""取り消し線のサンプル文書を作る (架空の作業要領書の改訂案)。

同じ内容を 4 つの形式で作る: Word、Excel、テキスト層のある PDF、スキャン PDF (画像だけの PDF)。
取り消し線のある箇所は EXPECT_STRUCK、紛らわしいが取り消し線ではない箇所 (下線、表の罫線) も入れてある。

usage:
  python samples/strike/make_sample.py samples/strike/source [--font <日本語の TrueType/OpenType フォント>]

python-docx、openpyxl、reportlab、pillow が要る。スキャン PDF は、文字を画像に描くのに日本語のフォントが要る。
"""
from __future__ import annotations

import argparse
import glob
import unicodedata
from pathlib import Path

TITLE = "空調設備 定期点検 作業要領書 (改訂案)"
# 段落。文字列の並びで、(文字列, 書式) の書式は "strike" (取り消し線) / "underline" (下線) / None
PARAS = [
    ("h", [("1. 適用範囲", None)]),
    ("p", [("本要領書は、第2棟の空調設備の定期点検に適用する。", None)]),
    ("p", [("第1棟の空調設備にも準用する。", "strike")]),
    ("h", [("2. 点検周期", None)]),
    ("p", [("点検は 6 か月ごとに実施する。", "strike")]),
    ("p", [("点検は 12 か月ごとに実施する。", None)]),
    ("h", [("3. 作業手順", None)]),
    ("p", [("(1) フィルターを取り外し、清掃する。", None)]),
    ("p", [("(2) ベルトの張りを", None), ("目視で確認する。", "strike"), ("張力計で測定し、記録する。", None)]),
    ("p", [("(3) 試運転を行い、", None), ("異音のないこと", "underline"), ("を確認する。", None)]),
    ("p", [("(4) 作業後、担当者が点検記録に押印する。", "strike")]),
    ("h", [("4. 判定基準", None)]),
]
# 表。セルは (文字列, 書式) の並び
TABLE = [
    [[("項目", None)], [("基準値", None)], [("備考", None)]],
    [[("風量", None)], [("1,200", "strike"), (" 1,500 m3/h 以上", None)], [("設計値の 90 % 以上", None)]],
    [[("運転電流", None)], [("12.5 A 以下", None)], [("定格の 110 % 以下", "strike")]],
    [[("振動", None)], [("異常のないこと", None)], [("触診による", None)]],
]
EXPECT_STRUCK = ["第1棟の空調設備にも準用する。", "点検は 6 か月ごとに実施する。", "目視で確認する。", "(4) 作業後、担当者が点検記録に押印する。", "1,200", "定格の 110 % 以下"]


def make_docx(path: Path):
    from docx import Document

    doc = Document()
    doc.add_heading(TITLE, 0)

    def runs(par, parts):
        for text, fmt in parts:
            run = par.add_run(text)
            run.font.strike = fmt == "strike"
            run.font.underline = fmt == "underline"

    for kind, parts in PARAS:
        if kind == "h":
            doc.add_heading(parts[0][0], 1)
        else:
            runs(doc.add_paragraph(), parts)
    table = doc.add_table(rows=len(TABLE), cols=3)
    table.style = "Table Grid"
    for r, row in enumerate(TABLE):
        for c, parts in enumerate(row):
            runs(table.cell(r, c).paragraphs[0], parts)
    doc.save(path)


def make_xlsx(path: Path):
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "判定基準"
    for r, row in enumerate(TABLE, start=1):
        for c, parts in enumerate(row, start=1):
            cell = ws.cell(r, c)
            if all(fmt == "strike" for _, fmt in parts):  # セル全体の書式
                cell.value, cell.font = "".join(t for t, _ in parts), Font(strike=True)
            elif any(fmt == "strike" for _, fmt in parts):  # セル内の一部の文字の書式
                cell.value = CellRichText(*[TextBlock(InlineFont(strike=True), t) if fmt == "strike" else t for t, fmt in parts])
            else:
                cell.value = "".join(t for t, _ in parts)
    ws2 = wb.create_sheet("作業手順")
    for r, (_, parts) in enumerate([p for p in PARAS if p[0] == "p"], start=1):
        cell = ws2.cell(r, 1)
        if all(fmt == "strike" for _, fmt in parts):
            cell.value, cell.font = "".join(t for t, _ in parts), Font(strike=True)
        elif any(fmt == "strike" for _, fmt in parts):
            cell.value = CellRichText(*[TextBlock(InlineFont(strike=True), t) if fmt == "strike" else t for t, fmt in parts])
        else:
            cell.value = "".join(t for t, _ in parts)
    wb.save(path)


class Page:
    """文字と線を置く場所を決める。PDF と画像で同じ配置にする。単位はポイント、原点は左上。"""

    W, H, LEFT, SIZE = 595, 842, 70, 11
    COLS = [70, 170, 340, 525]  # 表の縦の罫線

    def __init__(self, text_width):
        self.text_width = text_width  # (文字列, 大きさ) -> 幅
        self.texts: list[tuple[float, float, str, int]] = []  # (x, 上端 y, 文字列, 大きさ)
        self.lines: list[tuple[float, float, float, float]] = []  # 罫線・取り消し線・下線

    def parts(self, x: float, y: float, parts, size: int):
        for text, fmt in parts:
            w = self.text_width(text, size)
            self.texts.append((x, y, text, size))
            if fmt == "strike":
                self.lines.append((x, y + size * 0.52, x + w, y + size * 0.52))
            elif fmt == "underline":
                self.lines.append((x, y + size * 1.12, x + w, y + size * 1.12))
            x += w

    def layout(self):
        y = 60
        self.parts(self.LEFT, y, [(TITLE, None)], 16)
        y += 44
        for kind, parts in PARAS:
            if kind == "h":
                y += 8
                self.parts(self.LEFT, y, parts, 13)
                y += 26
            else:
                self.parts(self.LEFT + 12, y, parts, self.SIZE)
                y += 24
        top, row_h = y + 4, 28
        for r, row in enumerate(TABLE):
            for c, parts in enumerate(row):
                self.parts(self.COLS[c] + 8, top + r * row_h + 8, parts, self.SIZE)
        for r in range(len(TABLE) + 1):
            self.lines.append((self.COLS[0], top + r * row_h, self.COLS[-1], top + r * row_h))
        for x in self.COLS:
            self.lines.append((x, top, x, top + len(TABLE) * row_h))
        return self


def make_text_pdf(path: Path):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    font = "HeiseiKakuGo-W5"
    pdfmetrics.registerFont(UnicodeCIDFont(font))
    page = Page(lambda t, s: pdfmetrics.stringWidth(t, font, s)).layout()
    c = canvas.Canvas(str(path), pagesize=(Page.W, Page.H))
    for x, y, text, size in page.texts:
        c.setFont(font, size)
        c.drawString(x, Page.H - y - size * 0.88, text)
    c.setLineWidth(0.7)
    for x0, y0, x1, y1 in page.lines:
        c.line(x0, Page.H - y0, x1, Page.H - y1)
    c.save()


def make_scan_pdf(path: Path, font_path: str, dpi: int = 300):
    """ページを画像に描き、少しぼかして、画像だけの PDF にする (文字の情報を持たない)。"""
    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    k = dpi / 72
    fonts: dict[int, ImageFont.FreeTypeFont] = {}
    font = lambda size: fonts.setdefault(size, ImageFont.truetype(font_path, round(size * k)))
    page = Page(lambda t, s: font(s).getlength(t) / k).layout()
    img = Image.new("L", (round(Page.W * k), round(Page.H * k)), 255)
    d = ImageDraw.Draw(img)
    for x, y, text, size in page.texts:
        d.text((x * k, y * k), text, font=font(size), fill=0)
    for x0, y0, x1, y1 in page.lines:
        d.line((x0 * k, y0 * k, x1 * k, y1 * k), fill=0, width=3)
    img.filter(ImageFilter.GaussianBlur(0.8)).save(path, "PDF", resolution=dpi)


def find_font() -> str | None:
    # macOS のファイル名は濁点が分かれた形 (NFD) なので、正規化してから比べる
    for pattern, word in (("/System/Library/Fonts/*W3.ttc", "角ゴシック"), ("/usr/share/fonts/**/NotoSansCJK-Regular.ttc", ""), ("/usr/share/fonts/**/ipaexg.ttf", ""), ("C:/Windows/Fonts/meiryo.ttc", "")):
        hit = sorted(f for f in glob.glob(pattern, recursive=True) if word in unicodedata.normalize("NFC", f))
        if hit:
            return hit[0]
    return None


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("out_dir")
    ap.add_argument("--font", help="スキャン PDF に使う日本語のフォント (省略すると、よくある場所から探す)")
    a = ap.parse_args(argv)
    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "youryou.docx")
    make_xlsx(out / "youryou.xlsx")
    make_text_pdf(out / "youryou_text.pdf")
    font = a.font or find_font()
    if font is None:
        ap.error("日本語のフォントが見つかりません。--font で指定してください")
    make_scan_pdf(out / "youryou_scan.pdf", font)
    print(f"{out} に 4 つのファイルを作りました (スキャン PDF のフォント: {font})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
