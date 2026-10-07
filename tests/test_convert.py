import zipfile

import openpyxl
import pytest
from openpyxl.styles import Font

from dococr import convert, strike_docx
from helpers import legacy, make, needs_soffice, p, r

def test_modern_file_is_returned_as_is(tmp_path):
    f = tmp_path / "a.docx"
    f.write_bytes(b"")
    assert convert.to_modern(f, tmp_path / "out") == f
    assert convert.is_legacy("a.XLS") and not convert.is_legacy("a.xlsx")


@needs_soffice
def test_doc_keeps_tables_and_strike(tmp_path):
    grid = '<w:tblPr><w:tblW w:w="4000" w:type="dxa"/></w:tblPr><w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
    tbl = f"<w:tbl>{grid}<w:tr><w:tc>{p(r('圧力'))}</w:tc><w:tc>{p(r('5', '<w:strike/>'), r('7'))}</w:tc></w:tr></w:tbl>"
    doc = legacy(make(tmp_path, p(r("旧仕様", "<w:strike/>"), r("新仕様")) + tbl + p(r("末尾"))), "doc", tmp_path)
    out = convert.to_modern(doc, tmp_path / "out")
    assert out.suffix == ".docx"
    assert "<w:tbl>" in zipfile.ZipFile(out).read("word/document.xml").decode()  # 表が段落に崩れていない
    md, found = strike_docx.extract(doc)  # .doc を直接渡しても変換してから読む
    assert "~~旧仕様~~新仕様" in md and "| 圧力 | ~~5~~7 |" in md
    assert [f["text"] for f in found] == ["旧仕様", "5"]


@needs_soffice
def test_xls_keeps_merged_cells_hidden_columns_and_strike(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "作業計画"
    ws["A1"] = "資産番号"
    ws.merge_cells("B1:D1")
    ws["B1"] = "作業実績"
    ws["A2"] = "XYP1A"
    ws["B2"] = "●"
    ws["B2"].font = Font(strike=True)
    ws["C2"] = 2026
    ws.column_dimensions["C"].hidden = True
    src = tmp_path / "plan.xlsx"
    wb.save(src)
    out = convert.to_modern(legacy(src, "xls", tmp_path), tmp_path / "out")
    assert out.suffix == ".xlsx"
    ws2 = openpyxl.load_workbook(out)["作業計画"]
    assert [str(m) for m in ws2.merged_cells.ranges] == ["B1:D1"]
    assert ws2["A2"].value == "XYP1A" and ws2["C2"].value == 2026
    assert ws2["B2"].font.strike and not ws2["A2"].font.strike
    assert ws2.column_dimensions["C"].hidden
