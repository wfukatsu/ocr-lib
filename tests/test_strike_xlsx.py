import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font

from dococr import convert, strike_xlsx
from helpers import legacy, needs_soffice


def book(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "作業計画"
    ws["A1"] = "資産番号"
    ws["A2"] = "XYP1A"
    ws["B2"] = "旧周期 2Y"
    ws["B2"].font = Font(strike=True)  # セル全体
    ws["C2"] = CellRichText(TextBlock(InlineFont(strike=True), "部品交換"), "→外観確認")  # 一部の文字
    ws["D2"] = CellRichText("a|b ", TextBlock(InlineFont(strike=True), "~~x"))
    ws["E2"] = 2026
    ws["E2"].font = Font(strike=True)  # 数値
    ws["F2"] = "複数行\n二行目"
    ws["F2"].font = Font(strike=True)
    other = wb.create_sheet("実績")
    other["A1"] = "取り消しなし"
    hidden = wb.create_sheet("旧様式")
    hidden.sheet_state = "hidden"
    hidden["B3"] = "廃止"
    hidden["B3"].font = Font(strike=True)
    path = tmp_path / "plan.xlsx"
    wb.save(path)
    return path


def test_whole_cell_and_partial_strike(tmp_path):
    md, found = strike_xlsx.extract(book(tmp_path))
    by = {(f["sheet"], f["cell"]): f for f in found}
    assert by[("作業計画", "B2")] == {"sheet": "作業計画", "cell": "B2", "kind": "cell", "text": "旧周期 2Y", "struck": "旧周期 2Y", "md": "~~旧周期 2Y~~"}
    assert by[("作業計画", "C2")]["kind"] == "partial" and by[("作業計画", "C2")]["md"] == "~~部品交換~~→外観確認"
    assert by[("作業計画", "E2")]["md"] == "~~2026~~"
    assert ("作業計画", "A2") not in by and ("実績", "A1") not in by
    assert "## 作業計画\n\n| セル | 内容 |\n| --- | --- |\n| B2 | ~~旧周期 2Y~~ |" in md
    assert "## 実績" not in md


def test_escaping_line_breaks_and_hidden_sheets(tmp_path):
    md, found = strike_xlsx.extract(book(tmp_path))
    by = {(f["sheet"], f["cell"]): f for f in found}
    assert by[("作業計画", "D2")]["md"] == r"a\|b ~~\~\~x~~"  # 表の区切りと ~~ をエスケープする
    assert by[("作業計画", "F2")]["md"] == "~~複数行<br>二行目~~"
    assert "## 旧様式 (非表示シート)" in md and by[("旧様式", "B3")]["kind"] == "cell"


def test_run_with_its_own_font_does_not_inherit_the_cell_strike(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = CellRichText(TextBlock(InlineFont(strike=True), "旧"), TextBlock(InlineFont(sz=11), "新"), "素")
    ws["A1"].font = Font(strike=True)  # セルの書式は取り消し線。書式を持たない部分だけがこれに従う
    path = tmp_path / "a.xlsx"
    wb.save(path)
    assert strike_xlsx.extract(path)[1][0]["md"] == "~~旧~~新~~素~~"


def test_no_strike(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    path = tmp_path / "a.xlsx"
    wb.save(path)
    assert strike_xlsx.extract(path) == ("取り消し線付きのセルはありません。\n", [])


@needs_soffice
def test_xls_is_converted_first(tmp_path):
    _, found = strike_xlsx.extract(legacy(book(tmp_path), "xls", tmp_path))
    by = {(f["sheet"], f["cell"]): f["md"] for f in found}
    assert by[("作業計画", "B2")] == "~~旧周期 2Y~~" and by[("作業計画", "C2")] == "~~部品交換~~→外観確認"
